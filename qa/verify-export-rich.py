"""Read-only independent inspection of the rich XLSX regression fixture."""
import argparse
import io
import json
import warnings
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from PIL import Image


def verify(path: Path) -> dict:
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        assert len(names) == len(set(names)), "Duplicate archive parts"
        for name in names:
            if name.endswith((".xml", ".rels")):
                ET.fromstring(archive.read(name))
        media = [name for name in names if name.startswith("xl/media/")]
        assert media, "Expected embedded image"
        for name in media:
            with Image.open(io.BytesIO(archive.read(name))) as image:
                image.load()
                assert image.format == "PNG"
                assert image.width > 0 and image.height > 0
    with warnings.catch_warnings(record=True) as captured:
        warnings.simplefilter("always")
        workbook = openpyxl.load_workbook(path, data_only=False)
    assert not captured, [str(item.message) for item in captured]
    try:
        sheets = []
        bold_cells = 0
        images = 0
        for sheet in workbook:
            images += len(sheet._images)
            assert sheet.freeze_panes == "A2", "Header must remain frozen"
            assert all(cell.font.bold for cell in sheet[1]), "Header style missing"
            bold_cells += sum(bool(cell.font.bold) for row in sheet.iter_rows(min_row=2) for cell in row)
            for image in sheet._images:
                anchor = image.anchor._from
                assert 1 <= anchor.row < sheet.max_row, "Image outside data rows"
                assert 0 <= anchor.col < sheet.max_column, "Image outside columns"
            sheets.append({"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
                           "images": len(sheet._images)})
        assert images == len(media), "Images were not all resolved by the workbook reader"
        assert bold_cells, "Expected a styled data/total row"
        return {"sheets": sheets, "decoded_images": images, "bold_data_cells": bold_cells}
    finally:
        workbook.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    arguments = parser.parse_args()
    print(json.dumps(verify(arguments.workbook.resolve(strict=True)), ensure_ascii=False))
